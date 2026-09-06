#core.py
import json
import os
from pathlib import Path
from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel

load_dotenv()

client = Groq(api_key=os.getenv("GROQ_API_KEY"))
MODEL  = "openai/gpt-oss-120b"


# ── Pydantic Models ───────────────────────────────────────────────────────────

class JobD(BaseModel):
    role:                  str
    required_skills:       list[str]
    preferred_skills:      list[str]
    minimum_experience:    float | None
    education_requirements: list[str]
    responsibilities:      list[str]


class Experience(BaseModel):
    company:     str | None = None
    role:        str | None = None
    duration:    str | None = None
    description: str | None = None
    skills_used: list[str]  = []


class Resume(BaseModel):
    name:                   str | None = None
    email:                  str | None = None
    phone:                  str | None = None
    total_experience_years: float | None = None
    skills:                 list[str]       = []
    experiences:            list[Experience] = []
    education:              list[str]       = []
    projects:               list[str]       = []
    certifications:         list[str]       = []


class MatchResult(BaseModel):
    score:   float
    details: dict


# ── File Readers ──────────────────────────────────────────────────────────────

def read_pdf(file_path: str) -> str:
    from pypdf import PdfReader
    reader = PdfReader(file_path)
    text   = ""
    for page in reader.pages:
        chunk = page.extract_text()
        if chunk:
            text += chunk + "\n"
    return text


def read_docx(file_path: str) -> str:
    from docx import Document
    doc  = Document(file_path)
    text = ""
    for para in doc.paragraphs:
        if para.text.strip():
            text += para.text + "\n"
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"
    return text


def read_resume_file(file_path: str) -> str:
    path = Path(file_path)
    if path.suffix.lower() == ".pdf":
        return read_pdf(file_path)
    elif path.suffix.lower() == ".docx":
        return read_docx(file_path)
    raise ValueError(f"Unsupported file format: {path.suffix}")


# ── LLM Functions ─────────────────────────────────────────────────────────────

def analyze_job(jd_text: str) -> JobD:
    """Parse raw job description text → structured JobD object."""
    schema = JobD.model_json_schema()
    system = (
        "You are an expert HR assistant. "
        "Extract structured information from job descriptions. "
        f"Return ONLY valid JSON matching this schema: {schema} "
        "Do NOT return the schema itself — fill it with actual extracted data. "
        "Return null for minimum_experience if not mentioned. "
        "Return [] for any list with no data. Do not invent information."
    )
    resp = client.chat.completions.create(
        model    = MODEL,
        messages = [
            {"role": "system", "content": system},
            {"role": "user",   "content": f"Analyze this job description:\n{jd_text}"},
        ],
        response_format = {"type": "json_object"},
    )
    return JobD(**json.loads(resp.choices[0].message.content))


def parse_resume(resume_text: str) -> Resume:
    """Parse raw resume text → structured Resume object."""
    schema = Resume.model_json_schema()
    system = (
        "You are an expert resume parser. "
        "Extract information based on meaning, not just section headings. "
        "Internships count as experience. Extract skills from ALL sections. "
        f"Return ONLY valid JSON matching this schema: {schema} "
        "Return null for missing values, [] for empty lists. "
        "Do not invent information."
    )
    resp = client.chat.completions.create(
        model    = MODEL,
        messages = [
            {"role": "system", "content": system},
            {"role": "user",   "content": f"Parse this resume:\n{resume_text}"},
        ],
        response_format = {"type": "json_object"},
    )
    return Resume(**json.loads(resp.choices[0].message.content))


def get_match_score(job: JobD, resume: Resume) -> MatchResult:
    """Compare structured job vs resume → match score + details."""
    schema = MatchResult.model_json_schema()
    prompt = f"""
You are an experienced HR recruiter.

JOB DESCRIPTION:
{job.model_dump_json(indent=2)}

CANDIDATE RESUME:
{resume.model_dump_json(indent=2)}

Compare the candidate against the job. Return JSON matching: {schema}

The "details" dict MUST contain exactly these keys:
  - "candidate_name"   : string
  - "matching_skills"  : list of strings
  - "missing_skills"   : list of strings
  - "experience_met"   : boolean
  - "match_percentage" : float between 0 and 100
  - "verdict"          : 2-3 sentence plain-English summary for the recruiter

Set score equal to match_percentage.
"""
    resp = client.chat.completions.create(
        model    = MODEL,
        messages = [{"role": "user", "content": prompt}],
        response_format = {"type": "json_object"},
    )
    return MatchResult(**json.loads(resp.choices[0].message.content))
